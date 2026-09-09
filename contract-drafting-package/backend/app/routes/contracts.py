from fastapi import APIRouter, Depends, HTTPException, status, Query
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.core.deps import get_current_user
from app.models.user import User
from app.schemas.contract import (
    ContractCreate,
    ContractUpdate,
    ContractResponse,
    ContractListResponse,
    ContractGenerateRequest,
    ContractGenerateResponse,
    ContractSubmitApprovalRequest,
)
from app.services.contract_service import ContractService
from app.services.deepagents_service import get_deepagents_service
from app.services.workflow_service import (
    archive_contract,
    create_approval_workflow,
    save_workflow_run,
    workflow_to_dict,
)
from app.models.workflow import ContractWorkflowRun, WorkflowReport
from typing import Optional
import logging

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/contracts", tags=["合同管理"])


@router.post("/from-content", response_model=ContractResponse)
async def create_contract_from_content(
    payload: ContractCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """从粘贴的合同正文直接创建合同草稿（供审核结果页『提交审核』使用，不跑 AI 生成）。"""
    service = ContractService(db)
    if not payload.content or not payload.content.strip():
        raise HTTPException(status_code=422, detail="合同正文不能为空")
    contract = service.create_contract(payload, current_user.id)
    return contract


@router.post("/export-docx")
async def export_contract_docx(
    payload: dict,
    current_user: User = Depends(get_current_user),
):
    """用 python-docx 生成真正的 .docx 合同文件（解决 HTML 伪 .doc 下载后乱码的问题）。

    body: {"title": "合同标题", "content": "HTML 正文"}
    """
    import html as html_mod
    import io
    import re
    from urllib.parse import quote

    from docx import Document
    from fastapi.responses import Response

    title = str(payload.get("title") or "合同").strip() or "合同"
    content = payload.get("content") or ""

    # HTML → 纯文本段落（保留段落结构）
    text = re.sub(r"<br\s*/?>|</p>|</div>|</h[1-6]>|</li>|</tr>", "\n", content, flags=re.I)
    text = re.sub(r"<[^>]+>", "", text)
    text = html_mod.unescape(text)
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]

    doc = Document()
    doc.add_heading(title, level=0)
    for p in paragraphs:
        doc.add_paragraph(p)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    filename = f"{title}.docx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/generate", response_model=ContractGenerateResponse)
async def generate_contract(
    request: ContractGenerateRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    try:
        logger.info(f"Starting contract generation for user {current_user.username}")

        input_data = {
            "party_id": request.party_id,
            "customer_name": request.customer_name,
            "contract_type": request.contract_type,
            "contract_amount": float(request.amount) if request.amount else 0,
            "jurisdiction": request.jurisdiction or "中国",
            "industry": request.industry or "通用",
            "payment_terms": {},
            "delivery_schedule": {},
            "description": request.description,
            "requirements": request.requirements,
            "materials_text": request.materials_text,
            "materials": request.materials or [],
        }

        deepagents = get_deepagents_service()
        result = await deepagents.execute(
            f"drafting:{request.contract_type}:{current_user.id}",
            {
                "task_type": "drafting",
                "action": "generate_document_from_case",
                "context": input_data,
            },
        )

        if result.get("status") not in ("success", "completed"):
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=result.get("error", "合同生成失败")
            )

        contract_service = ContractService(db)
        generation_result = result.get("result", {})
        risk_result = generation_result.get("risk_assessment", {}) if isinstance(generation_result, dict) else {}
        compliance_result = generation_result.get("quality_check", {}) if isinstance(generation_result, dict) else {}
        similarity_result = generation_result.get("similarity_search", {}) if isinstance(generation_result, dict) else {}
        content_block = generation_result.get("content", {}) if isinstance(generation_result, dict) else {}
        generated_content = content_block.get("content") if isinstance(content_block, dict) else generation_result.get("content", "")
        generated_content = (generated_content or "").strip() if isinstance(generated_content, str) else str(generated_content or "").strip()
        if not generated_content:
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail="合同生成未返回正文，请检查主模型配置或重试",
            )

        contract_data = ContractCreate(
            title=f"{request.contract_type} - {request.customer_name or '未命名'}",
            contract_type=request.contract_type,
            customer_name=request.customer_name,
            amount=request.amount,
            content=generated_content,
        )

        contract = contract_service.create_contract(contract_data, current_user.id)
        workflow_run = save_workflow_run(
            db=db,
            contract=contract,
            user_id=current_user.id,
            request_payload=input_data,
            result=result,
        )
        approval_workflow = create_approval_workflow(
            db=db,
            contract=contract,
            submitted_by=current_user,
            summary=(result.get("contract_summary") if isinstance(result, dict) else None) or request.description or generated_content[:500],
            risk_score=risk_result.get("overall", {}).get("score") if isinstance(risk_result, dict) else None,
            risk_level=risk_result.get("overall", {}).get("level") if isinstance(risk_result, dict) else None,
        )
        db.commit()
        db.refresh(contract)
        db.refresh(approval_workflow)

        return ContractGenerateResponse(
            contract_id=contract.id,
            contract_number=contract.contract_number,
            content=generated_content,
            html_content=content_block.get("html_content", "") if isinstance(content_block, dict) else "",
            status=contract.status,
            risk_score=risk_result.get("overall", {}).get("score") if isinstance(risk_result, dict) else None,
            risk_level=risk_result.get("overall", {}).get("level") if isinstance(risk_result, dict) else None,
            similar_contracts=similarity_result.get("total_found", 0) if isinstance(similarity_result, dict) else 0,
            compliance_pass=True,
            risk_dimensions=risk_result if isinstance(risk_result, dict) else None,
            similar_contract_list=similarity_result.get("similar_contracts", []) if isinstance(similarity_result, dict) else None,
            compliance_issues=compliance_result.get("issues", []) if isinstance(compliance_result, dict) else None,
            compliance_summary={
                "status": compliance_result.get("status") if isinstance(compliance_result, dict) else None,
                "total_issues": compliance_result.get("total_issues") if isinstance(compliance_result, dict) else None,
                "high_severity": compliance_result.get("high_severity") if isinstance(compliance_result, dict) else None,
                "checked_rules": compliance_result.get("checked_rules", []) if isinstance(compliance_result, dict) else [],
            } if isinstance(compliance_result, dict) else None,
            contract_summary=result.get("contract_summary") if isinstance(result, dict) else None,
            data_collection=generation_result.get("data_collection") if isinstance(generation_result, dict) else None,
            regulation_matches=None,
            main_model_compliance=compliance_result if isinstance(compliance_result, dict) else None,
            template_selection=generation_result.get("template_selection") if isinstance(generation_result, dict) else None,
            workflow_stages=[
                {"key": "data_collection", "label": "数据采集", "status": "completed"},
                {"key": "similarity_search", "label": "本地向量检索", "status": "completed"},
                {"key": "compliance", "label": "法规匹配与合规检查", "status": "completed"},
                {"key": "risk", "label": "四维风险评估", "status": "completed"},
                {"key": "generation", "label": "模板与主模型生成", "status": "completed"},
                {"key": "summary", "label": "合同摘要确认", "status": "awaiting_confirmation"},
                {"key": "approval", "label": "提交审批", "status": "pending_confirmation"},
            ],
            workflow_id=approval_workflow.workflow_id,
            approval_workflow=workflow_to_dict(approval_workflow),
        )

    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"Contract generation error: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"合同生成失败: {str(e)}"
        )


@router.get("", response_model=ContractListResponse)
async def get_contracts(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    status_filter: Optional[str] = Query(None, alias="status"),
    contract_type: Optional[str] = None,
    customer_id: Optional[int] = None,
    search: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    service = ContractService(db)
    contracts, total = service.get_contracts(
        page=page,
        page_size=page_size,
        status=status_filter,
        contract_type=contract_type,
        search=search,
        user=current_user,
    )
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": contracts,
    }


@router.get("/{contract_id}", response_model=ContractResponse)
async def get_contract(
    contract_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    service = ContractService(db)
    contract = service.get_contract(contract_id)
    if not contract:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="合同不存在")
    return contract


@router.put("/{contract_id}", response_model=ContractResponse)
async def update_contract(
    contract_id: int,
    contract_data: ContractUpdate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    service = ContractService(db)
    contract = service.update_contract(contract_id, contract_data)
    if not contract:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="合同不存在")
    return contract


@router.delete("/{contract_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_contract(
    contract_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    service = ContractService(db)
    success = service.delete_contract(contract_id)
    if not success:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="合同不存在")
    return None


@router.post("/{contract_id}/submit-approval")
async def submit_contract_approval(
    contract_id: int,
    request: ContractSubmitApprovalRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    service = ContractService(db)
    contract = service.get_contract(contract_id)
    if not contract:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="合同不存在")

    workflow = create_approval_workflow(
        db=db,
        contract=contract,
        submitted_by=current_user,
        summary=request.summary,
        risk_score=request.risk_score,
        risk_level=request.risk_level,
    )
    db.commit()
    db.refresh(workflow)
    return {
        "status": "success",
        "workflow": workflow_to_dict(workflow),
        "contract": contract.to_dict(),
    }


@router.post("/{contract_id}/archive")
async def archive_contract_route(
    contract_id: int,
    payload: Optional[dict] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    service = ContractService(db)
    contract = service.get_contract(contract_id)
    if not contract:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="合同不存在")

    record = archive_contract(db, contract, current_user, (payload or {}).get("reason"))
    db.commit()
    db.refresh(contract)
    return {
        "status": "success",
        "archive_record_id": record.id,
        "contract": contract.to_dict(),
    }


@router.get("/{contract_id}/workflow")
async def get_contract_workflow(
    contract_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    contract = db.query(ContractWorkflowRun).filter(ContractWorkflowRun.contract_id == contract_id).order_by(ContractWorkflowRun.started_at.desc()).first()
    if not contract:
        return {"workflow": None, "reports": []}
    workflow = contract.workflow
    reports = db.query(WorkflowReport).filter(WorkflowReport.contract_id == contract_id).order_by(WorkflowReport.created_at.asc()).all()
    return {
        "workflow": workflow_to_dict(workflow) if workflow else None,
        "reports": [{
            "id": report.id,
            "report_type": report.report_type,
            "agent_name": report.agent_name,
            "status": report.status,
            "payload": report.payload,
            "created_at": report.created_at.isoformat() if report.created_at else None,
        } for report in reports],
    }


@router.get("/{contract_id}/reports")
async def get_contract_reports(
    contract_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    reports = db.query(WorkflowReport).filter(WorkflowReport.contract_id == contract_id).order_by(WorkflowReport.created_at.asc()).all()
    return {
        "total": len(reports),
        "items": [{
            "id": report.id,
            "report_type": report.report_type,
            "agent_name": report.agent_name,
            "status": report.status,
            "payload": report.payload,
            "created_at": report.created_at.isoformat() if report.created_at else None,
        } for report in reports],
    }

